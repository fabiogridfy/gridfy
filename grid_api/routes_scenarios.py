"""
grid_api/routes_scenarios.py
Endpoints para gestión y aplicación de escenarios de carga/generación.
"""
import os
import shutil

from fastapi import APIRouter, HTTPException, UploadFile, File, Form, Depends
from pydantic import BaseModel
from sqlalchemy.orm import Session

from db.session import get_db
from db import crud
from grid_network.scenario import apply_scenario, validate_scenario_excel
from grid_api.routes_network import get_net
from paths import resolve, to_relative, REDES_DIR

router = APIRouter()


def _scenario_folder(network_name: str) -> str:
    safe = network_name.replace(" ", "_").replace("/", "_")
    folder = os.path.join(REDES_DIR, safe, "escenarios")
    os.makedirs(folder, exist_ok=True)
    return folder


# ── GET /api/networks/{id}/scenarios ─────────────────────────────────────────
@router.get("/networks/{network_id}/scenarios")
def list_scenarios(network_id: int, db: Session = Depends(get_db)):
    scenarios = crud.list_scenarios(db, network_id)
    return [
        {
            "id":          s.id,
            "name":        s.name,
            "description": s.description,
            "is_active":   s.is_active,
            "created_at":  s.created_at.isoformat() if s.created_at else None,
        }
        for s in scenarios
    ]


# ── POST /api/networks/{id}/scenarios/upload ──────────────────────────────────
@router.post("/networks/{network_id}/scenarios/upload")
async def upload_scenario(
    network_id:  int,
    file:        UploadFile = File(...),
    name:        str        = Form(...),
    description: str        = Form(""),
    db:          Session    = Depends(get_db),
):
    db_net = crud.get_network(db, network_id)
    if not db_net:
        raise HTTPException(404, "Red no encontrada")

    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Solo se aceptan archivos Excel")

    # Guardar Excel
    try:
        folder     = _scenario_folder(db_net.name)
        safe_name  = name.replace(" ", "_").replace("/", "_")
        excel_path = os.path.join(folder, f"{safe_name}.xlsx")

        content = await file.read()
        with open(excel_path, "wb") as f:
            f.write(content)
    except Exception as e:
        raise HTTPException(500, f"Error guardando el archivo: {e}")

    # Validar estructura
    try:
        errors = validate_scenario_excel(excel_path)
    except Exception as e:
        if os.path.exists(excel_path): os.remove(excel_path)
        raise HTTPException(500, f"Error validando el Excel: {e}")

    if errors:
        os.remove(excel_path)
        raise HTTPException(422, {"errors": errors})

    # Registrar en BD (ruta relativa)
    try:
        scenario = crud.create_scenario(
            db, network_id=network_id, name=name,
            description=description, excel_path=to_relative(excel_path),
        )
    except Exception as e:
        raise HTTPException(500, f"Error registrando en BD: {e}")

    return {"id": scenario.id, "name": scenario.name}


# ── POST /api/networks/{id}/scenarios/{sid}/apply ─────────────────────────────
@router.post("/networks/{network_id}/scenarios/{scenario_id}/apply")
def apply_scenario_ep(network_id: int, scenario_id: int, db: Session = Depends(get_db)):
    """Aplica el escenario a la red activa (sobrescribe P/Q en memoria)."""
    from grid_api.routes_network import set_net
    from grid_network.manager import network_manager

    db_net = crud.get_network(db, network_id)
    if not db_net:
        raise HTTPException(404, "Red no encontrada")

    # If net not loaded or wrong network loaded, load it now
    net = get_net()
    if net is None or network_manager.network_id != network_id:
        try:
            network_manager.load_from_excel(resolve(db_net.excel_path), db_net.id, db_net.name)
            set_net(network_manager.net)
            net = get_net()
        except Exception as e:
            raise HTTPException(500, f"Error cargando la red: {e}")

    if net is None:
        raise HTTPException(503, "Red no disponible — cárgala primero desde el selector")

    scenario = crud.get_scenario(db, scenario_id)
    if not scenario or scenario.network_id != network_id:
        raise HTTPException(404, "Escenario no encontrado")

    result = apply_scenario(net, resolve(scenario.excel_path), voltage=db_net.voltage)

    # Marcar como activo
    crud.set_active_scenario(db, network_id, scenario_id)

    return {
        "ok":            True,
        "scenario_name": scenario.name,
        "updated_loads": result["updated_loads"],
        "updated_sgens": result["updated_sgens"],
        "not_found":     result["not_found"],
    }


# ── POST /api/networks/{id}/scenarios/reset ───────────────────────────────────
@router.post("/networks/{network_id}/scenarios/reset")
def reset_scenario(network_id: int, db: Session = Depends(get_db)):
    """Pone todas las cargas y generadores a 0 MW (estado vacío)."""
    net = get_net()
    net.load["p_mw"]   = 0.0
    net.load["q_mvar"] = 0.0
    net.sgen["p_mw"]   = 0.0
    net.sgen["q_mvar"] = 0.0
    crud.set_active_scenario(db, network_id, None)
    return {"ok": True, "message": "Escenario restablecido a 0 MW"}


# ── DELETE /api/networks/{id}/scenarios/{sid} ─────────────────────────────────
@router.delete("/networks/{network_id}/scenarios/{scenario_id}")
def delete_scenario(network_id: int, scenario_id: int, db: Session = Depends(get_db)):
    scenario = crud.get_scenario(db, scenario_id)
    if not scenario or scenario.network_id != network_id:
        raise HTTPException(404, "Escenario no encontrado")
    sc_excel = resolve(scenario.excel_path)
    if sc_excel and os.path.exists(sc_excel):
        import gc
        gc.collect()  # force close any open file handles
        try:
            os.remove(sc_excel)
        except PermissionError:
            import time
            time.sleep(0.5)
            os.remove(sc_excel)
    crud.delete_scenario(db, scenario_id)
    return {"ok": True}


# ── GET /api/networks/{id}/scenarios/download-current ─────────────────────────
from fastapi.responses import Response as _Resp
import io as _io

@router.get("/networks/{network_id}/scenarios/download-current")
def download_current_scenario(network_id: int, db=Depends(get_db)):
    """Genera un Excel descargable con el estado actual de cargas y generadores."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    net    = get_net()
    db_net = crud.get_network(db, network_id)
    if not db_net:
        raise HTTPException(404, "Red no encontrada")

    wb = openpyxl.Workbook()
    BRAND = "037A68"; WHITE = "FFFFFF"; LIGHT = "E8F5F2"

    def write_sheet(ws, df_elements, cols):
        ws.sheet_view.showGridLines = False
        for ci, col in enumerate(cols, 1):
            c = ws.cell(row=1, column=ci, value=col)
            c.font = Font(name="Arial", bold=True, color=WHITE, size=10)
            c.fill = PatternFill("solid", fgColor=BRAND)
            c.alignment = Alignment(horizontal="center")
        for ri, (_, row) in enumerate(df_elements.iterrows(), 2):
            bg = LIGHT if ri % 2 == 0 else "F5F6F8"
            for ci, col in enumerate(cols, 1):
                val = row.get(col, "")
                # Convert MW to kW
                if col in ("P (kW)", "Q (kVAR)"):
                    mw_col = "p_mw" if "P" in col else "q_mvar"
                    val = round(float(row.get(mw_col, 0)) * 1000, 4)
                elif col == "CUPS":
                    val = row.get("name", "")
                elif col == "Terminal":
                    bus_id = int(row.get("bus", 0))
                    val = net.bus.at[bus_id, "name"] if bus_id in net.bus.index else ""
                c = ws.cell(row=ri, column=ci, value=val)
                c.font = Font(name="Arial", size=9)
                c.fill = PatternFill("solid", fgColor=bg)

    # Cargas sheet
    ws_loads = wb.active; ws_loads.title = "Cargas"
    write_sheet(ws_loads, net.load, ["CUPS", "Terminal", "P (kW)", "Q (kVAR)"])

    # Generacion sheet
    ws_gen = wb.create_sheet("Generacion")
    write_sheet(ws_gen, net.sgen, ["CUPS", "Terminal", "P (kW)", "Q (kVAR)"])

    buf = _io.BytesIO()
    wb.save(buf); buf.seek(0)
    fname = f"Escenario_{db_net.name.replace(' ','_')}.xlsx"
    return _Resp(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'}
    )


# ── GET /api/networks/{id}/scenarios/template ─────────────────────────────────
@router.get("/networks/{network_id}/scenarios/template")
def download_scenario_template(network_id: int, db=Depends(get_db)):
    """
    Genera una plantilla Excel de escenario con los CUPS de la red activa
    y columnas P(kW) y Q(kVAR) vacías listas para rellenar.
    """
    import openpyxl, _io
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from fastapi.responses import Response as _R

    net    = get_net()
    db_net = crud.get_network(db, network_id)
    if not db_net:
        raise HTTPException(404, "Red no encontrada")

    BRAND = "037A68"; WHITE = "FFFFFF"; LIGHT = "E8F5F2"; MUTED = "F5F6F8"
    YELLOW_BG = "FEF9EC"; DARK = "130E1E"
    thin   = Side(style='thin', color="C8D0DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = openpyxl.Workbook()

    def make_sheet(ws, title, elements, element_type):
        ws.sheet_view.showGridLines = False

        # Title row
        ws.merge_cells("A1:D1")
        t = ws["A1"]
        t.value = f"Gridfy — Plantilla Escenario: {title} ({db_net.name})"
        t.font = Font(name='Arial', bold=True, size=12, color=WHITE)
        t.fill = PatternFill('solid', fgColor=DARK)
        t.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[1].height = 26

        # Instruction row
        ws.merge_cells("A2:D2")
        n = ws["A2"]
        n.value = "Rellena las columnas P (kW) y Q (kVAR). Deja en blanco los elementos que no quieras modificar."
        n.font = Font(name='Arial', size=9, color="6B6085", italic=True)
        n.fill = PatternFill('solid', fgColor=MUTED)
        n.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[2].height = 16

        # Headers
        headers = ["CUPS", "Terminal", "P (kW)", "Q (kVAR)"]
        widths   = [32, 52, 12, 12]
        for ci, (h, w) in enumerate(zip(headers, widths), 1):
            c = ws.cell(row=3, column=ci, value=h)
            c.font = Font(name='Arial', bold=True, color=WHITE, size=10)
            c.fill = PatternFill('solid', fgColor=BRAND)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = border
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.row_dimensions[3].height = 20

        # Data rows — CUPS and Terminal filled, P/Q empty
        for ri, (_, el) in enumerate(elements.iterrows(), 4):
            bg = LIGHT if ri % 2 == 0 else MUTED
            bus_id = int(el["bus"])
            terminal = net.bus.at[bus_id, "name"] if bus_id in net.bus.index else ""
            cups     = str(el.get("name", ""))

            for ci, val in enumerate([cups, terminal, "", ""], 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.font = Font(name='Arial', size=9,
                              color="6B6085" if ci > 2 else DARK)
                # P and Q columns in yellow to indicate they need filling
                c.fill = PatternFill('solid', fgColor=YELLOW_BG if ci > 2 else bg)
                c.alignment = Alignment(vertical='center',
                                        horizontal='right' if ci > 2 else 'left')
                c.border = border
            ws.row_dimensions[ri].height = 15

        # Freeze pane at row 4
        ws.freeze_panes = "A4"

    # Sheet 1: Cargas
    ws1 = wb.active
    ws1.title = "Cargas"
    make_sheet(ws1, "Cargas", net.load, "load")

    # Sheet 2: Generacion
    ws2 = wb.create_sheet("Generacion")
    make_sheet(ws2, "Generacion", net.sgen, "sgen")

    buf = _io.BytesIO()
    wb.save(buf); buf.seek(0)
    fname = f"Plantilla_Escenario_{db_net.name.replace(' ','_')}.xlsx"
    return _R(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'}
    )
