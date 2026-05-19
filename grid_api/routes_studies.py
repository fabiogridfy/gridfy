"""
routes_studies.py
Endpoints de estudios eléctricos: power flow, diagnóstico, creación de elementos.
"""
import pandapower as pp
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

from grid_network import run_powerflow, diagnose, get_bus_geo, safe_float
from grid_api.routes_network import get_net

router = APIRouter()


@router.post("/powerflow")
def run_pf():
    net = get_net()
    try:
        return run_powerflow(net)
    except ValueError as e:
        raise HTTPException(422, str(e))


@router.get("/diagnose")
def get_diagnose():
    net = get_net()
    return diagnose(net)


class CreateEl(BaseModel):
    bus: int
    name: str
    p_mw: float = 0.0
    q_mvar: float = 0.0

@router.post("/create/load")
def create_load(body: CreateEl):
    net = get_net()
    if body.bus not in net.bus.index:
        raise HTTPException(400, f"Bus {body.bus} no existe")
    idx = pp.create_load(net, bus=body.bus, name=body.name,
                         p_mw=body.p_mw, q_mvar=body.q_mvar)
    geo = get_bus_geo(net, body.bus)
    return {"id": int(idx), "bus": body.bus, "name": body.name,
            "p_mw": body.p_mw, "q_mvar": body.q_mvar,
            "lat": safe_float(geo["y"]) if geo else None,
            "lon": safe_float(geo["x"]) if geo else None}

@router.post("/create/sgen")
def create_sgen(body: CreateEl):
    net = get_net()
    if body.bus not in net.bus.index:
        raise HTTPException(400, f"Bus {body.bus} no existe")
    idx = pp.create_sgen(net, bus=body.bus, name=body.name,
                         p_mw=body.p_mw, q_mvar=body.q_mvar)
    geo = get_bus_geo(net, body.bus)
    return {"id": int(idx), "bus": body.bus, "name": body.name,
            "p_mw": body.p_mw, "q_mvar": body.q_mvar,
            "lat": safe_float(geo["y"]) if geo else None,
            "lon": safe_float(geo["x"]) if geo else None}


# ── POST /api/hc/map ──────────────────────────────────────────────────────────
from fastapi.responses import StreamingResponse as _SR
from pydantic import BaseModel as _BM

class HCMapRequest(_BM):
    mode:    str = "demand"      # demand | generation
    voltage: str = "BT"

@router.post("/hc/map")
def run_hc_map(body: HCMapRequest):
    """Calcula el Hosting Capacity de todos los buses. Progreso via SSE."""
    from grid_network.hc_map import hc_map_generator
    net = get_net()

    def generate():
        try:
            for msg in hc_map_generator(net, body.mode, body.voltage):
                yield f"data: {msg}\n\n"
        except Exception as e:
            yield f"data: error: {e}\n\n"
        finally:
            yield "data: done\n\n"

    return _SR(generate(), media_type="text/event-stream",
               headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ── POST /api/hc/download ─────────────────────────────────────────────────────
from fastapi.responses import Response as _HCResp
from pydantic import BaseModel as _HCBM
from typing import List as _List
import io as _hcio

class HCBusResult(_HCBM):
    bus_id:     int
    bus_name:   str
    hosting_kw: int
    saturated:  bool
    at_limit:   bool

class HCDownloadRequest(_HCBM):
    mode:     str
    voltage:  str
    p_max_kw: int
    buses:    _List[HCBusResult]

@router.post("/hc/download")
def download_hc_results(body: HCDownloadRequest):
    """Genera un Excel con los resultados del HC de toda la red."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hosting Capacity"
    ws.sheet_view.showGridLines = False

    BRAND = "037A68"; WHITE = "FFFFFF"; DARK = "130E1E"
    GREEN = "E8F5F2"; YELLOW = "FEF9EC"; RED = "FDE8E8"; MUTED = "F5F6F8"
    thin = Side(style='thin', color="C8D0DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws.merge_cells("A1:D1")
    t = ws["A1"]
    t.value = f"Gridfy — Hosting Capacity ({body.mode.upper()}) — {body.voltage}"
    t.font = Font(name='Arial', bold=True, size=13, color=WHITE)
    t.fill = PatternFill('solid', fgColor=DARK)
    t.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 28

    # Subtitle
    ws.merge_cells("A2:D2")
    s = ws["A2"]
    s.value = f"Limite de estudio: {body.p_max_kw:,} kW | Modo: {body.mode} | Tension: {body.voltage}"
    s.font = Font(name='Arial', size=9, color="6B6085", italic=True)
    s.fill = PatternFill('solid', fgColor=MUTED)
    s.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 16

    # Headers
    headers = ["Bus", "Tensión (kV)", "Capacidad disponible (kW)", "Limitación"]
    widths   = [55, 14, 18, 35]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=3, column=ci, value=h)
        c.font = Font(name='Arial', bold=True, color=WHITE, size=10)
        c.fill = PatternFill('solid', fgColor=BRAND)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[3].height = 20

    # Data rows — get voltage from active net if possible
    net = None
    try:
        from grid_api.routes_network import get_net
        net = get_net()
    except Exception:
        pass

    # Sort: lowest HC first (most constrained)
    sorted_buses = sorted(body.buses, key=lambda b: b.hosting_kw)

    for ri, b in enumerate(sorted_buses, 4):
        # Determine voltage
        vn_kv = "—"
        if net is not None and b.bus_id in net.bus.index:
            vn_kv = f"{net.bus.at[b.bus_id, 'vn_kv']:.3f}"

        # Limitation text
        if b.saturated:
            limit = "Red ya saturada sin nueva conexion"
            bg = RED
        elif b.at_limit:
            limit = f"HC >= {body.p_max_kw:,} kW (limite del estudio)"
            bg = GREEN
        elif b.hosting_kw == 0:
            limit = "Sin capacidad disponible"
            bg = RED
        elif b.hosting_kw < body.p_max_kw * 0.33:
            limit = "Capacidad limitada"
            bg = YELLOW
        else:
            limit = "Capacidad disponible"
            bg = GREEN

        row_vals = [b.bus_name, vn_kv, b.hosting_kw, limit]
        for ci, val in enumerate(row_vals, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font = Font(name='Arial', size=9,
                         bold=(ci == 3),
                         color="037A68" if bg == GREEN else ("C0392B" if bg == RED else "B45309"))
            c.fill = PatternFill('solid', fgColor=bg)
            c.alignment = Alignment(vertical='center',
                                    horizontal='right' if ci == 3 else 'left')
            c.border = border
        ws.row_dimensions[ri].height = 15

    # Summary row
    total = len(sorted_buses)
    admissible = sum(1 for b in sorted_buses if b.hosting_kw > 0)
    ws.merge_cells(f"A{total+4}:D{total+4}")
    summary = ws.cell(row=total+4, column=1,
        value=f"Total: {total} buses analizados | Admisibles: {admissible} | Sin capacidad: {total-admissible}")
    summary.font = Font(name='Arial', bold=True, size=9, color=WHITE)
    summary.fill = PatternFill('solid', fgColor=DARK)
    summary.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[total+4].height = 18

    buf = _hcio.BytesIO()
    wb.save(buf); buf.seek(0)
    from datetime import datetime
    fname = f"HC_{body.mode}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return _HCResp(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'}
    )


# ── POST /api/create/acometida ────────────────────────────────────────────────
from pydantic import BaseModel as _BMac

from typing import Optional as _OptAc
class CreateAcometidaBody(_BMac):
    name:        str
    p_mw:        float
    q_mvar:      float = 0.0
    is_gen:      bool  = False
    connect_bus: int
    conductor:   str   = "MANGUERA CU(4X6)"
    dist_m:      float = 10.0
    supply_lat:  _OptAc[float] = None
    supply_lon:  _OptAc[float] = None

@router.post("/create/acometida")
def create_acometida(body: CreateAcometidaBody):
    import pandapower as _pp
    from grid_network.conductor_db import get_conductor_params

    net = get_net()
    if body.connect_bus not in net.bus.index:
        raise HTTPException(400, f"Bus {body.connect_bus} no existe")

    vn_kv = float(net.bus.at[body.connect_bus, "vn_kv"])

    # Coordinates for new bus
    # pandapower bus_geodata uses x=lon, y=lat (NOT UTM)
    geo_x, geo_y = None, None
    if body.supply_lat and body.supply_lon:
        # Store as lon, lat directly
        geo_x = body.supply_lon  # x = longitude
        geo_y = body.supply_lat  # y = latitude
    if geo_x is None:
        geo = get_bus_geo(net, body.connect_bus)
        if geo:
            geo_x = safe_float(geo["x"])
            geo_y = safe_float(geo["y"])

    # 1. New bus at supply point
    new_bus_name = f"BUS_{body.name}"
    new_bus = _pp.create_bus(net, vn_kv=vn_kv, name=new_bus_name,
                             type="b",
                             geodata=(geo_x, geo_y) if geo_x else None)

    # 2. Acometida line
    cond = get_conductor_params(body.conductor)
    length_km = max(body.dist_m / 1000.0, 0.001)
    new_line = _pp.create_line_from_parameters(
        net,
        from_bus=new_bus, to_bus=body.connect_bus,
        length_km=length_km,
        r_ohm_per_km=cond["r_ohm_per_km"],
        x_ohm_per_km=cond["x_ohm_per_km"],
        c_nf_per_km=0.0,
        max_i_ka=cond["i_max_ka"],
        name=f"ACOM_{body.name}",
        in_service=True,
    )
    # Preservar el nombre del conductor y la sección para el Excel descargable
    net.line.at[new_line, "conductor"]   = cond.get("name") or body.conductor
    if cond.get("seccion_mm2") is not None:
        net.line.at[new_line, "seccion_mm2"] = float(cond["seccion_mm2"])

    # 3. Load or sgen
    if body.is_gen:
        el = _pp.create_sgen(net, bus=new_bus, name=body.name,
                             p_mw=body.p_mw, q_mvar=body.q_mvar)
        el_type = "sgen"
    else:
        el = _pp.create_load(net, bus=new_bus, name=body.name,
                             p_mw=body.p_mw, q_mvar=body.q_mvar)
        el_type = "load"

    return {
        "ok": True,
        "new_bus": int(new_bus), "new_line": int(new_line),
        "element": int(el), "el_type": el_type,
        "name": body.name, "bus_name": new_bus_name,
        "conductor": body.conductor, "dist_m": body.dist_m,
        "lat": body.supply_lat, "lon": body.supply_lon,
    }
