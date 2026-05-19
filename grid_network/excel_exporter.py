"""
grid_network/excel_exporter.py
Exporta una red pandapower al formato Excel canónico de Gridfy
(mismas hojas y columnas que el builder espera al cargar).
"""
import io
import math
import utm
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from grid_network.geodata import get_bus_geo


# Hoja → cabeceras esperadas por el builder
TERMINAL_COLS  = ["Nombre", "Tensión", "X", "Y"]
EXTGRID_COLS   = ["Nombre", "Terminal"]
LINEAS_COLS    = ["Nombre", "Terminal_i", "Terminal_j", "Longitud (km)",
                  "R/km", "X/km", "I max (kA)", "in_service"]
TRAFO_COLS     = ["code", "primary_node_code", "secondary_node_code",
                  "nominal_apparent_power_kVA", "Tension HV", "Tension LV",
                  "vk_percent", "vkr_percent", "pfe_kw", "i0_percent"]
LOADS_COLS     = ["CUPS", "Terminal", "P (MW)", "Q (MVAR)", "Pot. contratada (kW)"]
GEN_COLS       = ["CUPS", "Terminal", "P (MW)", "Q (MVAR)", "Pot. contratada (kW)"]


def _safe_float(v, default=0.0):
    try:
        f = float(v)
        if math.isnan(f) or math.isinf(f):
            return default
        return f
    except Exception:
        return default


def _bus_name(net, bus_idx) -> str:
    try:
        return str(net.bus.at[int(bus_idx), "name"])
    except Exception:
        return f"bus_{int(bus_idx)}"


def network_to_excel_bytes(net) -> bytes:
    """
    Serializa la red activa en un Excel en formato Gridfy.
    Devuelve los bytes del archivo (.xlsx).
    """
    # ── Terminal: convertir lon/lat (GeoJSON) → UTM (X, Y) ────────────────────
    rows_term = []
    for idx, row in net.bus.iterrows():
        geo = get_bus_geo(net, int(idx))
        if geo and geo["x"] is not None and geo["y"] is not None:
            try:
                x_utm, y_utm, _, _ = utm.from_latlon(geo["y"], geo["x"])
            except Exception:
                x_utm, y_utm = 0.0, 0.0
        else:
            x_utm, y_utm = 0.0, 0.0
        rows_term.append({
            "Nombre":  row["name"],
            "Tensión": _safe_float(row["vn_kv"]),
            "X":       round(x_utm, 3),
            "Y":       round(y_utm, 3),
        })
    df_term = pd.DataFrame(rows_term, columns=TERMINAL_COLS)

    # ── ExternalGrid ──────────────────────────────────────────────────────────
    rows_ext = []
    for idx, row in net.ext_grid.iterrows():
        rows_ext.append({
            "Nombre":   row.get("name", f"ext_{idx}"),
            "Terminal": _bus_name(net, row["bus"]),
        })
    df_ext = pd.DataFrame(rows_ext, columns=EXTGRID_COLS)

    # ── Líneas ────────────────────────────────────────────────────────────────
    rows_line = []
    for idx, row in net.line.iterrows():
        rows_line.append({
            "Nombre":         row.get("name", f"L_{idx}"),
            "Terminal_i":     _bus_name(net, row["from_bus"]),
            "Terminal_j":     _bus_name(net, row["to_bus"]),
            "Longitud (km)":  _safe_float(row["length_km"]),
            "R/km":           _safe_float(row["r_ohm_per_km"]),
            "X/km":           _safe_float(row["x_ohm_per_km"]),
            "I max (kA)":     _safe_float(row["max_i_ka"]),
            "in_service":     bool(row.get("in_service", True)),
        })
    df_lin = pd.DataFrame(rows_line, columns=LINEAS_COLS)

    # ── Transformadores ───────────────────────────────────────────────────────
    rows_tr = []
    for idx, row in net.trafo.iterrows():
        rows_tr.append({
            "code":                       row.get("name", f"TRF_{idx}"),
            "primary_node_code":          _bus_name(net, row["hv_bus"]),
            "secondary_node_code":        _bus_name(net, row["lv_bus"]),
            "nominal_apparent_power_kVA": _safe_float(row["sn_mva"]) * 1000.0,
            "Tension HV":                 _safe_float(row["vn_hv_kv"]),
            "Tension LV":                 _safe_float(row["vn_lv_kv"]),
            "vk_percent":                 _safe_float(row["vk_percent"]),
            "vkr_percent":                _safe_float(row["vkr_percent"]),
            "pfe_kw":                     _safe_float(row["pfe_kw"]),
            "i0_percent":                 _safe_float(row["i0_percent"]),
        })
    df_tr = pd.DataFrame(rows_tr, columns=TRAFO_COLS)

    # ── Loads_Data ────────────────────────────────────────────────────────────
    rows_ld = []
    for idx, row in net.load.iterrows():
        nom_kw = _safe_float(row.get("max_p_mw", 0)) * 1000.0
        rows_ld.append({
            "CUPS":                 row.get("name", f"load_{idx}"),
            "Terminal":             _bus_name(net, row["bus"]),
            "P (MW)":               _safe_float(row["p_mw"]),
            "Q (MVAR)":             _safe_float(row["q_mvar"]),
            "Pot. contratada (kW)": nom_kw,
        })
    df_ld = pd.DataFrame(rows_ld, columns=LOADS_COLS)

    # ── Gen_Data ──────────────────────────────────────────────────────────────
    rows_gn = []
    for idx, row in net.sgen.iterrows():
        nom_kw = _safe_float(row.get("max_p_mw", 0)) * 1000.0
        rows_gn.append({
            "CUPS":                 row.get("name", f"gen_{idx}"),
            "Terminal":             _bus_name(net, row["bus"]),
            "P (MW)":               _safe_float(row["p_mw"]),
            "Q (MVAR)":             _safe_float(row["q_mvar"]),
            "Pot. contratada (kW)": nom_kw,
        })
    df_gn = pd.DataFrame(rows_gn, columns=GEN_COLS)

    # ── Escribir Excel con estilo Gridfy ──────────────────────────────────────
    BRAND = "037A68"; WHITE = "FFFFFF"; LIGHT = "E8F5F2"; MUTED = "F5F6F8"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def write_sheet(name: str, df: pd.DataFrame):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        for ci, col in enumerate(df.columns, 1):
            c = ws.cell(row=1, column=ci, value=col)
            c.font = Font(name="Arial", bold=True, color=WHITE, size=10)
            c.fill = PatternFill("solid", fgColor=BRAND)
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(ci)].width = max(18, len(str(col)) + 4)
            ws.row_dimensions[1].height = 20
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            bg = LIGHT if ri % 2 == 0 else MUTED
            for ci, val in enumerate(row, 1):
                cell_val = None if (val is None or (isinstance(val, float) and (math.isnan(val) or math.isinf(val)))) else val
                c = ws.cell(row=ri, column=ci, value=cell_val)
                c.font = Font(name="Arial", size=9)
                c.fill = PatternFill("solid", fgColor=bg)
                c.alignment = Alignment(vertical="center")
            ws.row_dimensions[ri].height = 15

    write_sheet("Terminal",        df_term)
    write_sheet("ExternalGrid",    df_ext)
    write_sheet("Lineas",          df_lin)
    write_sheet("Transformadores", df_tr)
    if not df_ld.empty:
        write_sheet("Loads_Data",  df_ld)
    if not df_gn.empty:
        write_sheet("Gen_Data",    df_gn)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
